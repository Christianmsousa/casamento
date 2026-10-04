#!/usr/bin/env python3
"""
Converte planilha/convidados-whatsapp.json → planilha/convites-assessoria-vip.xlsx
usando o template planilha/convites-base.xlsx (Assessoria VIP).
"""

from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
JSON_PATH = ROOT / "planilha" / "convidados-whatsapp.json"
TEMPLATE_PATH = ROOT / "planilha" / "convites-base.xlsx"
OUTPUT_PATH = ROOT / "planilha" / "convites-assessoria-vip.xlsx"

SHEET_NAME = "Planilha de convidados"
DATA_START_ROW = 3


def group_for(inv: dict) -> str:
    if inv.get("grupo_convite"):
        return str(inv["grupo_convite"])
    tipo = inv.get("tipo_convite")
    if tipo == "Padrinho":
        return "Padrinhos"
    if tipo == "Cerimonial":
        return "Cerimonial"
    if inv.get("responsavel_envio") == "julia":
        return "Amigos da Noiva"
    return "Amigos do Noivo"


def faixa_vip(pessoa: dict) -> str:
    tipo = pessoa.get("tipo")
    fe = (pessoa.get("faixa_etaria") or "").lower()

    if (
        "recém" in fe
        or "recem" in fe
        or "menor de 3" in fe
        or re.search(r"\b(1|2)\s*anos?\b", fe)
    ):
        return "Criança de colo"
    if "mais de 12" in fe or "aproximadamente 12" in fe:
        return "Adolescente"
    if tipo == "Criança" or fe:
        return "Criança"
    return "Adulto"


def custo_vip(faixa: str) -> str:
    return {
        "Adulto": "Inteira",
        "Idoso": "Inteira",
        "Adolescente": "Inteira",
        "Criança": "Meia",
        "Criança de colo": "Gratuita",
    }.get(faixa, "Inteira")


def situacao_vip(inv: dict, pessoa: dict) -> str:
    return (
        pessoa.get("situacao")
        or inv.get("situacao")
        or "Pendente"
    )


def phone_parts(inv: dict) -> tuple[str, str]:
    tels = inv.get("telefones") or []
    if not tels:
        return "", ""
    n = tels[0].get("numero_internacional") or ""
    digits = re.sub(r"\D", "", n)
    if digits.startswith("55") and len(digits) >= 12:
        rest = digits[2:]
        ddd, num = rest[:2], rest[2:]
        if len(num) == 9:
            fmt = f"{ddd} {num[:5]}-{num[5:]}"
        elif len(num) == 8:
            fmt = f"{ddd} {num[:4]}-{num[4:]}"
        else:
            fmt = f"{ddd} {num}"
        return "55", fmt
    if digits.startswith("351"):
        return "351", digits[3:]
    return "", n


def nome_pessoa(pessoa: dict) -> str:
    if pessoa.get("nome"):
        return str(pessoa["nome"])
    ref = pessoa.get("referencia_original") or pessoa.get("faixa_etaria") or "convidado"
    return str(ref).strip().title()


def observacao(inv: dict) -> str:
    parts = list(inv.get("pendencias") or [])
    if not (inv.get("telefones") or []):
        if "telefone_ausente" not in parts:
            parts.append("telefone_ausente")
    return "; ".join(dict.fromkeys(parts))


def iter_invites(data: dict):
    for side in ("christian", "julia", "definir_responsavel"):
        for inv in data["listas"].get(side, []):
            yield inv


def build_rows(data: dict) -> list[dict]:
    rows: list[dict] = []
    for inv in iter_invites(data):
        ddi, tel = phone_parts(inv)
        grp = group_for(inv)
        obs = observacao(inv)
        people = inv.get("convidados") or []
        for i, pessoa in enumerate(people):
            faixa = faixa_vip(pessoa)
            rows.append(
                {
                    "A": inv.get("nome_convite", "") if i == 0 else "",
                    "B": ddi if i == 0 and ddi else "",
                    "C": tel if i == 0 and tel else "",
                    "D": grp if i == 0 else "",
                    "E": obs if i == 0 and obs else "",
                    "F": nome_pessoa(pessoa),
                    "G": pessoa.get("sexo") or "",
                    "H": faixa,
                    "I": custo_vip(faixa),
                    "J": situacao_vip(inv, pessoa),
                    "K": "",
                }
            )
    return rows


def clear_data_rows(ws) -> None:
    max_row = max(ws.max_row, DATA_START_ROW)
    for r in range(DATA_START_ROW, max_row + 1):
        for c in range(1, 12):
            ws.cell(r, c).value = None


def write_rows(ws, rows: list[dict]) -> None:
    for offset, row in enumerate(rows):
        r = DATA_START_ROW + offset
        for col, key in enumerate("ABCDEFGHIJK", start=1):
            ws.cell(r, col).value = row[key] if row[key] != "" else None


def main() -> None:
    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    rows = build_rows(data)

    shutil.copy2(TEMPLATE_PATH, OUTPUT_PATH)
    wb = load_workbook(OUTPUT_PATH)
    ws = wb[SHEET_NAME]
    clear_data_rows(ws)
    write_rows(ws, rows)
    wb.save(OUTPUT_PATH)

    invites = list(iter_invites(data))
    stats_faixa = Counter(r["H"] for r in rows)
    stats_grupo = Counter(r["D"] for r in rows if r["D"])
    stats_sit = Counter(r["J"] for r in rows)
    no_phone = [inv["nome_convite"] for inv in invites if not inv.get("telefones")]
    placeholders = [
        r["F"]
        for r in rows
        if r["F"]
        in {"Esposa Diogo", "Filho Diogo", "Filha Adriano"}
        or r["F"].startswith("(")
    ]

    print(f"Gerado: {OUTPUT_PATH}")
    print(f"Convites: {len(invites)} | Pessoas: {len(rows)}")
    print(f"Faixas: {dict(stats_faixa)}")
    print(f"Grupos: {dict(stats_grupo)}")
    print(f"Situação: {dict(stats_sit)}")
    print(f"Sem telefone: {no_phone or 'nenhum'}")
    print(f"Placeholders: {placeholders}")


if __name__ == "__main__":
    main()
